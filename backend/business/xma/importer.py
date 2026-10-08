from __future__ import annotations

import hashlib
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

import openpyxl

from .schema import Closure, Scene, SourceRef, Transfer, XmaAircraft, XmaDataset, XmaFlight

BEIJING = timezone(timedelta(hours=8))
DEFAULT_WORKBOOK = Path(__file__).resolve().parents[3] / "data/厦航-天池比赛数据/data/XMAE_data.xlsx"
HEADERS = {
    "航班": ["航班ID", "日期", "国际/国内", "航班号", "起飞机场", "降落机场", "起飞时间", "降落时间", "飞机ID", "机型", "旅客数", "联程旅客数", "座位数", "重要系数"],
    "航线-飞机限制": ["起飞机场", "降落机场", "飞机ID"],
    "机场关闭限制": ["机场", "关闭时间", "开放时间", "生效日期", "失效日期"],
    "台风场景": ["开始时间", "结束时间", "影响类型", "机场", "航班ID", "飞机", "停机数"],
    "飞行时间": ["飞机机型", "起飞机场", "降落机场", "飞行时间（分钟）"],
    "机场": ["机场", "国内/国际"],
    "中转时间限制": ["进港航班ID", "出港航班ID", "最短转机时限（分钟）", "中转旅客人数"],
}


def identifier(value) -> str:
    if value is None:
        raise ValueError("missing identifier")
    return str(int(value)) if isinstance(value, (int, float)) and int(value) == value else str(value).strip()


def timestamp(value) -> datetime:
    if isinstance(value, str):
        value = datetime.fromisoformat(value)
    if not isinstance(value, datetime):
        raise ValueError(f"invalid timestamp: {value!r}")
    return value.replace(tzinfo=BEIJING) if value.tzinfo is None else value.astimezone(BEIJING)


def minute_of_day(value) -> int:
    if isinstance(value, str):
        parts = value.split(":")
        return int(parts[0]) * 60 + int(parts[1])
    return value.hour * 60 + value.minute


def load_workbook(source: Path | bytes = DEFAULT_WORKBOOK) -> XmaDataset:
    raw = source if isinstance(source, bytes) else Path(source).read_bytes()
    workbook = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True)
    tables = {}
    try:
        for sheet, headers in HEADERS.items():
            if sheet not in workbook.sheetnames:
                raise ValueError(f"missing sheet: {sheet}")
            rows = list(workbook[sheet].values)
            actual = [str(x).strip() for x in rows[0][:len(headers)]]
            if actual != headers:
                raise ValueError(f"invalid headers in {sheet}: {actual}")
            tables[sheet] = [(i, r) for i, r in enumerate(rows[1:], 2) if any(x is not None for x in r)]
    finally:
        workbook.close()
    flights = []
    for i, r in tables["航班"]:
        flights.append(XmaFlight(flight_id=identifier(r[0]), date=timestamp(r[1]).date().isoformat(),
            domestic=r[2] == "国内", flight_no=identifier(r[3]), origin=identifier(r[4]), destination=identifier(r[5]),
            sched_dep=timestamp(r[6]), sched_arr=timestamp(r[7]), original_aircraft=identifier(r[8]),
            equipment=identifier(r[9]), passengers=r[10], through_passengers=r[11], seats=r[12], importance=r[13],
            source=SourceRef(sheet="航班", row=i)))
    rotations = defaultdict(list)
    for f in flights:
        rotations[f.original_aircraft].append(f)
    aircraft, connected, turns = [], [], {}
    for tail, rotation in sorted(rotations.items()):
        rotation.sort(key=lambda f: (f.sched_dep, f.flight_id))
        if len({(f.equipment, f.seats) for f in rotation}) != 1:
            raise ValueError(f"inconsistent equipment or seats for tail {tail}")
        aircraft.append(XmaAircraft(tail_id=tail, equipment=rotation[0].equipment, seats=rotation[0].seats,
            initial_station=rotation[0].origin, original_terminal=rotation[-1].destination,
            original_rotation=[f.flight_id for f in rotation]))
        for a, b in zip(rotation, rotation[1:]):
            gap = int((b.sched_dep-a.sched_arr).total_seconds()/60)
            if gap < 0 or a.destination != b.origin:
                raise ValueError(f"broken original rotation: {a.flight_id} -> {b.flight_id}")
            turns[f"{a.flight_id}#{b.flight_id}"] = gap
            if a.date == b.date and a.flight_no == b.flight_no:
                if a.through_passengers != b.through_passengers:
                    raise ValueError("inconsistent connected passenger count")
                connected.append((a.flight_id, b.flight_id))
    scenes = [Scene(start=timestamp(r[0]), end=timestamp(r[1]), kind={"降落":"arrival", "起飞":"departure", "停机":"parking"}[r[2]],
        airport=identifier(r[3]), flight_id=identifier(r[4]) if r[4] is not None else None,
        tail_id=identifier(r[5]) if r[5] is not None else None, parking_limit=int(r[6] or 0) if r[2]=="停机" else None,
        source=SourceRef(sheet="台风场景",row=i)) for i,r in tables["台风场景"]]
    data = XmaDataset(dataset_id="xma_2017_"+hashlib.sha256(raw).hexdigest()[:12], source_sha256=hashlib.sha256(raw).hexdigest(),
        adjustment_start=timestamp("2017-05-06T06:00:00"), adjustment_end=timestamp("2017-05-09T00:00:00"),
        flights=flights, aircraft=aircraft, airports={identifier(r[0]): bool(r[1]) for _,r in tables["机场"]},
        route_prohibitions=[tuple(identifier(x) for x in r[:3]) for _,r in tables["航线-飞机限制"]],
        closures=[Closure(airport=identifier(r[0]),close_minutes=minute_of_day(r[1]),open_minutes=minute_of_day(r[2]),
            effective_date=timestamp(r[3]).date().isoformat(),expiration_date=timestamp(r[4]).date().isoformat(),
            source=SourceRef(sheet="机场关闭限制",row=i)) for i,r in tables["机场关闭限制"]], scenes=scenes,
        travel_minutes={"#".join(identifier(x) for x in r[:3]):int(r[3]) for _,r in tables["飞行时间"]},
        transfers=[Transfer(inbound=identifier(r[0]),outbound=identifier(r[1]),minimum_minutes=r[2],count=r[3],
            source=SourceRef(sheet="中转时间限制",row=i)) for i,r in tables["中转时间限制"]],
        connected_pairs=connected, original_turn_minutes=turns,
        movement_windows=[(timestamp("2017-05-06T15:00:00"),timestamp("2017-05-06T16:00:00")),
            (timestamp("2017-05-07T17:00:00"),timestamp("2017-05-07T19:00:00"))],
        notes=["Source preserved. Crew and maintenance absent and disabled.",
            "Adjustment window and movement limits sourced from bundled InputData.java, not workbook cells.",
            "Scene/closure boundaries are open; movement windows and adjustment window are inclusive.",
            "Parking means a consecutive-flight ground interval spans the entire scene, per Scene.java."])
    stationary=sum(f.origin==f.destination for f in flights)
    if stationary:
        data.notes.append(f"Preserved {stationary} same-airport flight records; no undocumented maintenance interpretation is imposed.")
    data.source_tables={sheet:[{"row":i,"values":{h:timestamp(v).isoformat() if isinstance(v,datetime) else str(v) if hasattr(v,"hour") else v
        for h,v in zip(HEADERS[sheet],r)}} for i,r in rows] for sheet,rows in tables.items()}
    duplicates=len(tables["飞行时间"])-len(data.travel_minutes)
    if duplicates:
        data.notes.append(f"Travel-time table has {duplicates} duplicate keys; original rows retained. Effective lookup follows InputData.java last-row-wins semantics.")
    errors = validate_dataset(data)
    if errors:
        raise ValueError("; ".join(errors))
    return data


def validate_dataset(data: XmaDataset) -> list[str]:
    errors = []
    flights = {f.flight_id:f for f in data.flights}
    tails = {a.tail_id:a for a in data.aircraft}
    if not flights or len(flights) != len(data.flights) or len(tails) != len(data.aircraft):
        errors.append("empty flights or duplicate flight/tail IDs")
    if data.adjustment_start>=data.adjustment_end:
        errors.append("invalid adjustment window")
    covered=[]
    for tail,a in tails.items():
        rotation=[flights[fid] for fid in a.original_rotation if fid in flights]
        covered.extend(a.original_rotation)
        if not rotation or len(rotation)!=len(a.original_rotation):
            errors.append(f"missing or empty original rotation on {tail}")
            continue
        if rotation[0].origin!=a.initial_station or rotation[-1].destination!=a.original_terminal:
            errors.append(f"original boundary station mismatch on {tail}")
        for f in rotation:
            if (f.original_aircraft,f.equipment,f.seats)!=(tail,a.equipment,a.seats):
                errors.append(f"original tail/equipment/seats mismatch on {f.flight_id}")
        for left,right in zip(rotation,rotation[1:]):
            gap=int((right.sched_dep-left.sched_arr).total_seconds()/60)
            if gap<0 or left.destination!=right.origin:
                errors.append(f"broken original rotation: {left.flight_id}/{right.flight_id}")
            if data.original_turn_minutes.get(f"{left.flight_id}#{right.flight_id}")!=gap:
                errors.append(f"original turn exception mismatch: {left.flight_id}/{right.flight_id}")
    if Counter(covered)!=Counter(flights.keys()):
        errors.append("original rotations must cover each flight exactly once")
    for left,right in data.connected_pairs:
        if left not in flights or right not in flights:
            errors.append(f"missing connected flight: {left}/{right}")
        elif flights[left].original_aircraft!=flights[right].original_aircraft or flights[left].through_passengers!=flights[right].through_passengers:
            errors.append(f"inconsistent connected pair: {left}/{right}")
        elif f"{left}#{right}" not in data.original_turn_minutes:
            errors.append(f"connected pair is not originally consecutive: {left}/{right}")
    if any(s.airport not in data.airports for s in data.scenes) or any(c.airport not in data.airports for c in data.closures):
        errors.append("unknown airport in source restrictions")
    if any(start.tzinfo is None or end.tzinfo is None or start>=end for start,end in data.movement_windows):
        errors.append("invalid movement window")
    incoming, outgoing = Counter(), Counter()
    transfer_keys=[(t.inbound,t.outbound) for t in data.transfers]
    if len(transfer_keys)!=len(set(transfer_keys)):
        errors.append("duplicate transfer relationships must be resolved using source last-row-wins semantics")
    for t in data.transfers:
        if t.inbound not in flights and t.inbound not in data.boundary_flights or t.outbound not in flights:
            errors.append(f"transfer reference missing: {t.inbound}/{t.outbound}")
            continue
        a,b=flights.get(t.inbound) or data.boundary_flights[t.inbound],flights[t.outbound]
        if a.destination != b.origin:
            errors.append(f"transfer station mismatch: {t.inbound}/{t.outbound}")
        incoming[t.outbound] += t.count
        outgoing[t.inbound] += t.count
    pair_members = {f for p in data.connected_pairs for f in p}
    if len(pair_members) != 2*len(data.connected_pairs):
        errors.append("connected pairs overlap")
    for f in data.flights:
        if f.original_aircraft not in tails or f.origin not in data.airports or f.destination not in data.airports:
            errors.append(f"unknown resource/station on flight {f.flight_id}")
        if incoming[f.flight_id]+outgoing[f.flight_id] > min(f.passengers, f.seats-f.through_passengers):
            errors.append(f"transfer passenger accounting invalid on {f.flight_id}")
        if f.through_passengers and f.flight_id not in pair_members:
            errors.append(f"through passenger missing connected pair on {f.flight_id}")
    return errors


def select_subset(data: XmaDataset, target_flights: int | None) -> XmaDataset:
    """Whole-tail subinstances, conditioned on original inbound boundary flights.

    They are separate research problems, not feasible full-workbook solutions.
    No transfer into a selected flight is discarded or synthesized.
    """
    if target_flights is None or target_flights>=len(data.flights):
        return data
    selected=[]
    count=0
    # Stable airport/tail ordering keeps studies reproducible and includes typhoon tails first.
    affected={s.airport for s in data.scenes if s.kind in ("arrival","departure")}
    flights={f.flight_id:f for f in data.flights}
    ordered=sorted(data.aircraft,key=lambda a:(not any(flights[fid].origin in affected or flights[fid].destination in affected
        for fid in a.original_rotation),a.tail_id))
    for a in ordered:
        selected.append(a)
        count+=len(a.original_rotation)
        if count>=target_flights:
            break
    ids={fid for a in selected for fid in a.original_rotation}
    transfers=[t for t in data.transfers if t.outbound in ids]
    boundary={t.inbound:flights[t.inbound] for t in transfers if t.inbound not in ids}
    return data.model_copy(update={"dataset_id":data.dataset_id+f"_tails_{len(selected)}","flights":[f for f in data.flights if f.flight_id in ids],
        "aircraft":selected,"transfers":transfers,"boundary_flights":boundary,
        "connected_pairs":[p for p in data.connected_pairs if p[0] in ids],
        "notes":[*data.notes,f"Whole-tail subset: {len(ids)} flights. {len(boundary)} inbound boundary flights frozen at their original times. Outbound effects outside the subset are excluded from its objective."]})


def preview(data: XmaDataset) -> dict:
    return {"dataset_id":data.dataset_id,"source_sha256":data.source_sha256,"counts":{
        "flights":len(data.flights),"aircraft":len(data.aircraft),"airports":len(data.airports),
        "equipment_types":len({a.equipment for a in data.aircraft}),"route_prohibitions":len(data.route_prohibitions),
        "closures":len(data.closures),"scenes":len(data.scenes),"travel_times":len(data.source_tables.get("飞行时间",data.travel_minutes)),
        "transfers":len(data.transfers),"connected_pairs":len(data.connected_pairs),
        "oversold_flights":sum(f.passengers+f.through_passengers>f.seats for f in data.flights),
        "short_original_turns":sum(v<50 for v in data.original_turn_minutes.values())},
        "planning_start":min(f.sched_dep for f in data.flights).isoformat(),
        "original_planning_end":max(f.sched_arr for f in data.flights).isoformat(),
        "adjustment_start":data.adjustment_start.isoformat(),"adjustment_end":data.adjustment_end.isoformat(),
        "issues":validate_dataset(data),"notes":data.notes,
        "passenger_semantics":"Counts are flight-leg loads; their sum is not unique travelers."}
